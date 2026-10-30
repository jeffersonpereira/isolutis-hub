"""Geração de planilha Excel para relatórios financeiros (DRE e Fluxo de Caixa)."""

import base64
import io

from app.documents.reportlab_base import Identidade

NAVY = "0B2D3B"
TEAL = "4F6B8F"
SOFT = "EEF1F6"
WHITE = "FFFFFF"
TEXTO_SUAVE = "56647A"
OK = "4A9D6F"
BAD = "A85D52"


def relatorio_xlsx(
    dre: list[dict],
    fluxo: list[dict],
    empresa: Identidade,
    ano: int,
) -> bytes:
    """Gera arquivo Excel com duas abas: DRE e Fluxo de Caixa.

    Args:
        dre: Dados do DRE (saída de ``calcular_dre``).
        fluxo: Dados do Fluxo de Caixa (saída de ``calcular_fluxo_caixa``).
        empresa: Identidade de quem emite o relatório (nome, logo, dados fiscais).
        ano: Ano de referência.

    Returns:
        Conteúdo do arquivo .xlsx como bytes.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook(write_only=False)

    # --- Aba DRE ---
    ws_dre = wb.active
    ws_dre.title = "DRE"

    _cabecalho_planilha(ws_dre, empresa, f"{empresa.nome} — DRE {ano}", n_colunas=4)
    cabecalhos_dre = ["Mês", "Receitas (R$)", "Custos (R$)", "Resultado (R$)"]
    _escrever_cabecalho(ws_dre, cabecalhos_dre, cor_fundo=NAVY, cor_fonte=WHITE)
    linha_cabecalho_dre = ws_dre.max_row

    for row in dre:
        resultado = row["resultado"]
        ws_dre.append([row["nome_mes"], row["receitas"], row["custos"], resultado])
        ultima = ws_dre.max_row
        for col in range(1, 5):
            cell = ws_dre.cell(row=ultima, column=col)
            cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
            if col > 1:
                cell.number_format = '#,##0.00'
        # Colorir resultado positivo/negativo
        cel_resultado = ws_dre.cell(row=ultima, column=4)
        if isinstance(resultado, (int, float)):
            cel_resultado.font = Font(
                color=OK if resultado >= 0 else BAD,
                bold=resultado == 0,
            )

    linha_final_dados_dre = ws_dre.max_row
    _ajustar_colunas(ws_dre, [20, 18, 18, 18])

    # Linha de totais
    total_receitas = sum(r["receitas"] for r in dre)
    total_custos = sum(r["custos"] for r in dre)
    total_resultado = total_receitas - total_custos
    ws_dre.append(["Total", total_receitas, total_custos, total_resultado])
    ultima = ws_dre.max_row
    for col in range(1, 5):
        cell = ws_dre.cell(row=ultima, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=SOFT)
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
        if col > 1:
            cell.number_format = '#,##0.00'

    if dre:
        _grafico_mensal(
            ws_dre, "Receitas x Custos", linha_cabecalho_dre, linha_final_dados_dre,
            col_categoria=1, col_min=2, col_max=3, ancora="F2",
        )  # fmt: skip

    # --- Aba Fluxo de Caixa ---
    ws_fc = wb.create_sheet("Fluxo de Caixa")

    _cabecalho_planilha(ws_fc, empresa, f"{empresa.nome} — Fluxo de Caixa {ano}", n_colunas=5)
    cabecalhos_fc = ["Mês", "Entradas (R$)", "Saídas (R$)", "Saldo do Mês (R$)", "Saldo Acumulado (R$)"]
    _escrever_cabecalho(ws_fc, cabecalhos_fc, cor_fundo=NAVY, cor_fonte=WHITE)
    linha_cabecalho_fc = ws_fc.max_row

    for row in fluxo:
        ws_fc.append([
            row["nome_mes"],
            row["entradas"],
            row["saidas"],
            row["saldo_mes"],
            row["saldo_acumulado"],
        ])
        ultima = ws_fc.max_row
        for col in range(1, 6):
            cell = ws_fc.cell(row=ultima, column=col)
            cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
            if col > 1:
                cell.number_format = '#,##0.00'
        # Colorir saldo do mês e saldo acumulado
        for col_idx, key in ((4, "saldo_mes"), (5, "saldo_acumulado")):
            val = row[key]
            if isinstance(val, (int, float)):
                ws_fc.cell(row=ultima, column=col_idx).font = Font(color=OK if val >= 0 else BAD)

    linha_final_dados_fc = ws_fc.max_row
    _ajustar_colunas(ws_fc, [20, 18, 18, 20, 22])

    if fluxo:
        _grafico_mensal(
            ws_fc, "Entradas x Saídas", linha_cabecalho_fc, linha_final_dados_fc,
            col_categoria=1, col_min=2, col_max=3, ancora="G2",
        )  # fmt: skip

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _cabecalho_planilha(ws: object, empresa: Identidade, titulo: str, n_colunas: int) -> None:
    """Escreve o título (nome da empresa + relatório/ano), a linha fiscal (documento/endereço/contato,
    quando houver) e embute a logomarca, se houver e for uma imagem válida."""
    from openpyxl.styles import Alignment, Font

    ws.append([titulo])
    cell = ws.cell(row=ws.max_row, column=1)
    cell.font = Font(bold=True, size=13)
    cell.alignment = Alignment(horizontal="left")

    linha_fiscal = " · ".join(p for p in (empresa.documento, empresa.endereco, empresa.contato) if p)
    if linha_fiscal:
        ws.append([linha_fiscal])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = Font(size=9, color=TEXTO_SUAVE)
        cell.alignment = Alignment(horizontal="left")

    ws.append([])  # linha em branco

    imagem = _logo_imagem(empresa.logo)
    if imagem:
        from openpyxl.utils import get_column_letter

        ws.add_image(imagem, f"{get_column_letter(n_colunas + 2)}1")


def _logo_imagem(logo_data_url: str | None) -> object | None:
    """Decodifica a logomarca para uma imagem do openpyxl, ou ``None`` se não houver logo ou se os
    bytes armazenados não formarem uma imagem válida (o cadastro só confere a assinatura do arquivo,
    não o conteúdo completo) — nesse caso a planilha fica sem logo, em vez de falhar a exportação."""
    if not logo_data_url or "base64," not in logo_data_url:
        return None
    from openpyxl.drawing.image import Image

    _, b64 = logo_data_url.split("base64,", 1)
    try:
        imagem = Image(io.BytesIO(base64.b64decode(b64)))
        if imagem.width > 140:
            escala = 140 / imagem.width
            imagem.width, imagem.height = 140, imagem.height * escala
        return imagem
    except Exception:
        return None


def _grafico_mensal(
    ws: object, titulo: str, linha_cabecalho: int, linha_final: int,
    col_categoria: int, col_min: int, col_max: int, ancora: str,
) -> None:  # fmt: skip
    """Gráfico de barras nativo do Excel (editável) com a evolução mensal das duas séries principais."""
    from openpyxl.chart import BarChart, Reference

    grafico = BarChart()
    grafico.type = "col"
    grafico.title = titulo
    grafico.height = 7.5
    grafico.width = 16
    grafico.legend.position = "b"
    grafico.y_axis.majorGridlines = None

    dados = Reference(ws, min_col=col_min, max_col=col_max, min_row=linha_cabecalho, max_row=linha_final)
    categorias = Reference(ws, min_col=col_categoria, min_row=linha_cabecalho + 1, max_row=linha_final)
    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)
    ws.add_chart(grafico, ancora)


def _escrever_cabecalho(ws: object, colunas: list[str], cor_fundo: str, cor_fonte: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.append(colunas)
    linha = ws.max_row
    for col in range(1, len(colunas) + 1):
        cell = ws.cell(row=linha, column=col)
        cell.font = Font(bold=True, color=cor_fonte)
        cell.fill = PatternFill("solid", fgColor=cor_fundo)
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")


def _ajustar_colunas(ws: object, larguras: list[int]) -> None:
    from openpyxl.utils import get_column_letter

    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
